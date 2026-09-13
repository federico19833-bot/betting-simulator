#!/usr/bin/env python3
"""Build LivePick signal workbook, lookup JSON and filterable HTML data."""

from __future__ import annotations

import csv
import json
import math
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "livepick_sorgente.csv"
OUT_XLSX = ROOT / "docs" / "livepick_segnali.xlsx"
OUT_PARTITE = ROOT / "docs" / "livepick_partite.json"
OUT_LOOKUP = ROOT / "docs" / "livepick_lookup.json"
OUT_CSV = ROOT / "data" / "livepick_partite.csv"

MULTIWORD_COUNTRIES = [
    "Antigua & Barbuda",
    "Bosnia & Herzegovina",
    "Burkina Faso",
    "Costa Rica",
    "Czech Republic",
    "Dominican Republic",
    "El Salvador",
    "Faroe Islands",
    "Hong Kong",
    "Ivory Coast",
    "New Zealand",
    "Northern Ireland",
    "North Macedonia",
    "Papua New Guinea",
    "Republic of Ireland",
    "Saint Kitts & Nevis",
    "San Marino",
    "Saudi Arabia",
    "Sierra Leone",
    "Solomon Islands",
    "South Africa",
    "South Korea",
    "South East Asian",
    "South Asian",
    "Sri Lanka",
    "Trinidad & Tobago",
    "United Arab Emirates",
    "West Asian",
]

ALIASES = {
    "Holland": "Netherlands",
    "Türkiye": "Turkey",
    "Czechia": "Czech Republic",
    "UAE": "United Arab Emirates",
    "USA": "USA",
    "Macedonia": "North Macedonia",
    "CAMBODIA_PREM_L": "Cambodia",
}

SPECIAL_PREFIXES = [
    ("Campeonato Brasileiro", "Brazil"),
    ("Brazilian Matches", "Brazil"),
    ("Copa do Brasil", "Brazil"),
    ("Copa Libertadores", "Sudamerica"),
    ("Copa Sudamericana", "Sudamerica"),
    ("Copa Peru", "Peru"),
    ("Copa Internacional", "Internazionale"),
    ("Coppa Italia", "Italy"),
    ("Europe Friendlies", "Europa (amichevoli)"),
    ("Elite Club Friendlies", "Elite (amichevoli)"),
    ("World Club Friendlies", "Mondo (amichevoli)"),
    ("Club Friendly", "Amichevoli"),
    ("Friendly Match", "Amichevoli"),
    ("Friendlies", "Amichevoli"),
    ("International Match", "Nazionali"),
    ("U19 International", "Nazionali U19"),
    ("U20 International", "Nazionali U20"),
    ("U23 International", "Nazionali U23"),
    ("U19 Tournament", "Nazionali U19"),
    ("U20 Tournament", "Nazionali U20"),
    ("Youth Tournament", "Giovanili internazionali"),
    ("UEFA", "UEFA"),
    ("AFC", "AFC"),
    ("CAF", "CAF"),
    ("CONCACAF", "CONCACAF"),
    ("FIFA", "FIFA"),
    ("Euro ", "Europa (nazionali)"),
    ("European", "Europa (nazionali)"),
    ("Africa Cup", "Africa (nazionali)"),
    ("Asia - World Cup", "Asia (nazionali)"),
    ("AFF Cup", "Asia (nazionali)"),
    ("Arab Club", "Arabo"),
    ("Cosafa", "Africa (nazionali)"),
    ("Premier League International Cup", "Inghilterra (internazionale)"),
    ("West Asian", "Asia (nazionali)"),
    ("South Asian", "Asia (nazionali)"),
    ("South East Asian", "Asia (nazionali)"),
]

P_FROM_MINUTES = [60, 63, 64, 65, 69, 75, 80, 81, 85]


def parse_score(value: str | None) -> tuple[int, int] | None:
    if not value:
        return None
    parts = str(value).replace("–", "-").split("-")
    if len(parts) != 2:
        return None
    try:
        return int(parts[0].strip()), int(parts[1].strip())
    except ValueError:
        return None


def extract_paese(league: str) -> str:
    raw = (league or "").strip()
    if not raw:
        return "Sconosciuto"
    if raw in ALIASES:
        return ALIASES[raw]
    for prefix, paese in SPECIAL_PREFIXES:
        if raw.startswith(prefix):
            return paese
    for name in MULTIWORD_COUNTRIES:
        if raw.startswith(name):
            return ALIASES.get(name, name)
    first = raw.split()[0]
    return ALIASES.get(first, first)


def wilson(k: int, n: int, z: float = 1.96) -> tuple[float, float, float]:
    if n <= 0:
        return 0.0, 0.0, 0.0
    p = k / n
    den = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / den
    err = z * math.sqrt((p * (1 - p) + z * z / (4 * n)) / n) / den
    return p, max(0.0, centre - err), min(1.0, centre + err)


def poisson_ge1(lam: float) -> float:
    if lam <= 0:
        return 0.0
    return 1.0 - math.exp(-lam)


def load_rows() -> list[dict]:
    rows: list[dict] = []
    with SRC.open(newline="", encoding="utf-8") as fh:
        for i, raw in enumerate(csv.DictReader(fh), start=1):
            ht = parse_score(raw.get("Half-Time Result"))
            ft = parse_score(raw.get("Full-Time Result"))
            pk = parse_score(raw.get("Picked Score")) or ht
            try:
                minute = int(raw.get("Picked Minute") or "")
            except ValueError:
                minute = None
            if not (ht and ft and pk and minute is not None):
                continue
            gol_dopo = (ft[0] + ft[1]) - (pk[0] + pk[1])
            casa_dopo = ft[0] - pk[0]
            ospite_dopo = ft[1] - pk[1]
            if gol_dopo < 0 or casa_dopo < 0 or ospite_dopo < 0:
                continue
            scarto = pk[0] - pk[1]
            if scarto > 0:
                stato = "Vantaggio casa"
            elif scarto < 0:
                stato = "Vantaggio ospite"
            else:
                stato = "Pareggio"
            residui = max(90 - minute, 0)
            rows.append(
                {
                    "id": i,
                    "paese": extract_paese(raw.get("League") or ""),
                    "campionato": (raw.get("League") or "").strip(),
                    "casa": (raw.get("Home Team Name") or "").strip(),
                    "ospite": (raw.get("Away Team Name") or "").strip(),
                    "ht": f"{ht[0]}-{ht[1]}",
                    "ht_casa": ht[0],
                    "ht_ospite": ht[1],
                    "segnale_minuto": minute,
                    "segnale_risultato": f"{pk[0]}-{pk[1]}",
                    "segnale_casa": pk[0],
                    "segnale_ospite": pk[1],
                    "segnale_totale": pk[0] + pk[1],
                    "segnale_scarto": abs(scarto),
                    "segnale_stato": stato,
                    "ft": f"{ft[0]}-{ft[1]}",
                    "ft_casa": ft[0],
                    "ft_ospite": ft[1],
                    "ft_totale": ft[0] + ft[1],
                    "gol_dopo": gol_dopo,
                    "altro_gol": gol_dopo > 0,
                    "casa_dopo": casa_dopo,
                    "ospite_dopo": ospite_dopo,
                    "minuti_residui_90": residui,
                }
            )
    return rows


def agg(rows: list[dict]) -> dict:
    n = len(rows)
    if n == 0:
        return {
            "n": 0,
            "altro_gol": 0,
            "p_altro_gol": 0.0,
            "ic95_low": 0.0,
            "ic95_high": 0.0,
            "e_gol_dopo": 0.0,
            "p_0": 0.0,
            "p_1": 0.0,
            "p_2plus": 0.0,
            "lambda_min": 0.0,
            "avg_minuto": 0.0,
        }
    k = sum(1 for r in rows if r["altro_gol"])
    p, lo, hi = wilson(k, n)
    e = sum(r["gol_dopo"] for r in rows) / n
    tot_min = sum(r["minuti_residui_90"] for r in rows)
    tot_gol = sum(r["gol_dopo"] for r in rows)
    lam = (tot_gol / tot_min) if tot_min else 0.0
    return {
        "n": n,
        "altro_gol": k,
        "p_altro_gol": round(p, 4),
        "ic95_low": round(lo, 4),
        "ic95_high": round(hi, 4),
        "e_gol_dopo": round(e, 4),
        "p_0": round(sum(1 for r in rows if r["gol_dopo"] <= 0) / n, 4),
        "p_1": round(sum(1 for r in rows if r["gol_dopo"] == 1) / n, 4),
        "p_2plus": round(sum(1 for r in rows if r["gol_dopo"] >= 2) / n, 4),
        "lambda_min": round(lam, 6),
        "avg_minuto": round(sum(r["segnale_minuto"] for r in rows) / n, 2),
    }


def p_from_minute(lam: float) -> dict:
    out = {}
    for m in P_FROM_MINUTES:
        remain = max(90 - m, 0)
        out[str(m)] = {
            "residui_90": remain,
            "e_gol_90": round(lam * remain, 4),
            "p_90": round(poisson_ge1(lam * remain), 4),
            "e_gol_90_3": round(lam * (remain + 3), 4),
            "p_90_3": round(poisson_ge1(lam * (remain + 3)), 4),
        }
    return out


def group(rows: list[dict], key) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        out[str(key(r))].append(r)
    return dict(out)


def build_lookup(rows: list[dict]) -> dict:
    late = [r for r in rows if r["segnale_minuto"] >= 60]
    globale = agg(rows)
    globale_late = agg(late)
    lookup = {
        "meta": {
            "n_partite": len(rows),
            "minuto_segnale_min": min(r["segnale_minuto"] for r in rows),
            "minuto_segnale_max": max(r["segnale_minuto"] for r in rows),
            "generato": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
            "nota": (
                "p_altro_gol = frequenza empirica di almeno un gol dal minuto del segnale al FT. "
                "p_da_minuto usa Poisson sul ritmo gol (lambda_min) del sottoinsieme tardi (>=60') "
                "se disponibile, altrimenti sul totale."
            ),
        },
        "globale": {
            **globale,
            "lambda_min_tardi": globale_late["lambda_min"],
            "p_da_minuto": p_from_minute(globale_late["lambda_min"] or globale["lambda_min"]),
            "per_minuto": {k: agg(v) for k, v in sorted(group(rows, lambda r: r["segnale_minuto"]).items(), key=lambda x: int(x[0]))},
            "per_risultato": {k: agg(v) for k, v in sorted(group(rows, lambda r: r["segnale_risultato"]).items())},
            "per_stato": {k: agg(v) for k, v in group(rows, lambda r: r["segnale_stato"]).items()},
        },
        "paesi": {},
        "campionati": {},
    }

    for paese, xs in sorted(group(rows, lambda r: r["paese"]).items()):
        late_p = [r for r in xs if r["segnale_minuto"] >= 60]
        lam = (agg(late_p)["lambda_min"] if late_p else 0.0) or agg(xs)["lambda_min"]
        lookup["paesi"][paese] = {
            **agg(xs),
            "lambda_min_tardi": agg(late_p)["lambda_min"] if late_p else agg(xs)["lambda_min"],
            "p_da_minuto": p_from_minute(lam),
            "per_minuto": {k: agg(v) for k, v in sorted(group(xs, lambda r: r["segnale_minuto"]).items(), key=lambda x: int(x[0]))},
            "per_risultato": {k: agg(v) for k, v in sorted(group(xs, lambda r: r["segnale_risultato"]).items())},
            "per_minuto_risultato": {
                f"{r0['segnale_minuto']}|{r0['segnale_risultato']}": agg(v)
                for key, v in group(xs, lambda r: (r["segnale_minuto"], r["segnale_risultato"])).items()
                for r0 in [v[0]]
            },
            "campionati": {k: agg(v) for k, v in sorted(group(xs, lambda r: r["campionato"]).items())},
        }

    for camp, xs in sorted(group(rows, lambda r: r["campionato"]).items()):
        lookup["campionati"][camp] = {
            **agg(xs),
            "paese": xs[0]["paese"],
            "per_minuto": {k: agg(v) for k, v in sorted(group(xs, lambda r: r["segnale_minuto"]).items(), key=lambda x: int(x[0]))},
        }
    return lookup


HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(bold=True, color="F8FAFC", name="Calibri")
SI_FILL = PatternFill("solid", fgColor="DCFCE7")
NO_FILL = PatternFill("solid", fgColor="FEE2E2")
THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
PCT = "0.0%"
NUM = "0.00"


def style_header(ws, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32


def autosize(ws, min_w=10, max_w=36):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min_w
        for cell in col[:80]:
            if cell.value is None:
                continue
            width = max(width, min(max_w, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def write_sheet(ws, headers: list[str], data: list[list], pct_cols: set[int] | None = None, num_cols: set[int] | None = None):
    pct_cols = pct_cols or set()
    num_cols = num_cols or set()
    ws.append(headers)
    for row in data:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.border = THIN
            cell.alignment = Alignment(horizontal="center" if c > 1 else "left")
            if c in pct_cols and isinstance(cell.value, (int, float)):
                cell.number_format = PCT
            if c in num_cols and isinstance(cell.value, (int, float)):
                cell.number_format = NUM
    style_header(ws, len(headers))
    autosize(ws)


def build_xlsx(rows: list[dict], lookup: dict) -> None:
    wb = Workbook()

    ws0 = wb.active
    ws0.title = "Come_usare"
    ws0["A1"] = "LivePick — foglio segnali di ingresso"
    ws0["A1"].font = Font(bold=True, size=16, color="0F172A")
    notes = [
        "",
        "Come leggere il foglio",
        "1) Partite: ogni riga è un segnale LivePick (minuto + risultato di ingresso) e l'esito a fine partita.",
        "2) Altro_gol = SI se dal segnale al 90' è uscito almeno un gol in più.",
        "3) Per_minuto / Per_paese / Per_risultato: probabilità empiriche già calcolate.",
        "4) P_da_minuto: stima Poisson di un gol DA un minuto successivo (es. 69' o 81'), non dal segnale.",
        "",
        "Domande tipo da fare all'agente",
        "- India, segnale al 64', risultato 2-1: probabilità di un altro gol?",
        "- Shillong Premier League, dal segnale.",
        "- Al 81', quanto resta rispetto al segnale del 64'?",
        "",
        f"Partite: {len(rows)}",
        f"Minuto segnale: {lookup['meta']['minuto_segnale_min']}–{lookup['meta']['minuto_segnale_max']}",
        f"P globale altro gol dal segnale: {lookup['globale']['p_altro_gol']:.1%} ({lookup['globale']['altro_gol']}/{lookup['globale']['n']})",
        f"Generato: {lookup['meta']['generato']}",
        "",
        "File collegati: docs/livepick.html (filtri), docs/livepick_lookup.json (risposte precalcolate).",
    ]
    for i, line in enumerate(notes, start=1):
        ws0[f"A{i}"] = line
        ws0[f"A{i}"].font = Font(bold=i in (1, 3, 10), size=12 if i == 1 else 11)
    ws0.column_dimensions["A"].width = 110

    headers_p = [
        "ID",
        "Paese",
        "Campionato",
        "Casa",
        "Ospite",
        "HT",
        "Segnale_minuto",
        "Segnale_risultato",
        "Segnale_totale",
        "Segnale_scarto",
        "Segnale_stato",
        "FT",
        "Gol_dopo_segnale",
        "Altro_gol",
        "Casa_dopo",
        "Ospite_dopo",
        "Minuti_residui_90",
    ]
    data_p = [
        [
            r["id"],
            r["paese"],
            r["campionato"],
            r["casa"],
            r["ospite"],
            r["ht"],
            r["segnale_minuto"],
            r["segnale_risultato"],
            r["segnale_totale"],
            r["segnale_scarto"],
            r["segnale_stato"],
            r["ft"],
            r["gol_dopo"],
            "SI" if r["altro_gol"] else "NO",
            r["casa_dopo"],
            r["ospite_dopo"],
            r["minuti_residui_90"],
        ]
        for r in rows
    ]
    ws = wb.create_sheet("Partite")
    write_sheet(ws, headers_p, data_p)
    last_row = ws.max_row
    ws.conditional_formatting.add(
        f"N2:N{last_row}",
        FormulaRule(formula=['N2="SI"'], fill=SI_FILL),
    )
    ws.conditional_formatting.add(
        f"N2:N{last_row}",
        FormulaRule(formula=['N2="NO"'], fill=NO_FILL),
    )
    tab = Table(displayName="PartiteSegnale", ref=f"A1:{get_column_letter(len(headers_p))}{last_row}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)

    def rows_from_agg(items: list[tuple[str, dict]], extra_headers: list[str] | None = None, extra_vals=None):
        headers = (extra_headers or []) + [
            "Chiave",
            "N",
            "Altro_gol_n",
            "P_altro_gol",
            "IC95_low",
            "IC95_high",
            "E_gol_dopo",
            "P_0",
            "P_1",
            "P_2plus",
            "Lambda_min",
            "Avg_minuto_segnale",
        ]
        data = []
        for key, a in items:
            prefix = extra_vals(key, a) if extra_vals else []
            data.append(
                prefix
                + [
                    key,
                    a["n"],
                    a["altro_gol"],
                    a["p_altro_gol"],
                    a["ic95_low"],
                    a["ic95_high"],
                    a["e_gol_dopo"],
                    a["p_0"],
                    a["p_1"],
                    a["p_2plus"],
                    a["lambda_min"],
                    a["avg_minuto"],
                ]
            )
        return headers, data

    h, d = rows_from_agg(list(lookup["globale"]["per_minuto"].items()))
    ws = wb.create_sheet("Per_minuto")
    write_sheet(ws, h, d, pct_cols={4, 5, 6, 8, 9, 10}, num_cols={7, 11, 12})
    ws["A1"] = "Minuto_segnale"
    ws.cell(1, 1).fill = HEADER_FILL
    ws.cell(1, 1).font = HEADER_FONT

    paese_items = [(p, a) for p, a in lookup["paesi"].items()]
    paese_items.sort(key=lambda x: -x[1]["n"])
    h, d = rows_from_agg(paese_items)
    ws = wb.create_sheet("Per_paese")
    write_sheet(ws, h, d, pct_cols={4, 5, 6, 8, 9, 10}, num_cols={7, 11, 12})
    ws["A1"] = "Paese"
    ws.cell(1, 1).fill = HEADER_FILL
    ws.cell(1, 1).font = HEADER_FONT

    pm = []
    for paese, pdata in lookup["paesi"].items():
        for minuto, a in pdata["per_minuto"].items():
            if a["n"] >= 1:
                pm.append((paese, minuto, a))
    pm.sort(key=lambda x: (x[0], int(x[1])))
    headers = [
        "Paese",
        "Minuto_segnale",
        "N",
        "Altro_gol_n",
        "P_altro_gol",
        "IC95_low",
        "IC95_high",
        "E_gol_dopo",
        "P_0",
        "P_1",
        "P_2plus",
    ]
    data = [
        [p, int(m), a["n"], a["altro_gol"], a["p_altro_gol"], a["ic95_low"], a["ic95_high"], a["e_gol_dopo"], a["p_0"], a["p_1"], a["p_2plus"]]
        for p, m, a in pm
    ]
    ws = wb.create_sheet("Per_paese_minuto")
    write_sheet(ws, headers, data, pct_cols={5, 6, 7, 9, 10, 11}, num_cols={8})

    h, d = rows_from_agg(sorted(lookup["globale"]["per_risultato"].items(), key=lambda x: -x[1]["n"]))
    ws = wb.create_sheet("Per_risultato")
    write_sheet(ws, h, d, pct_cols={4, 5, 6, 8, 9, 10}, num_cols={7, 11, 12})
    ws["A1"] = "Risultato_segnale"
    ws.cell(1, 1).fill = HEADER_FILL
    ws.cell(1, 1).font = HEADER_FONT

    pr = []
    for paese, pdata in lookup["paesi"].items():
        for score, a in pdata["per_risultato"].items():
            if a["n"] >= 1:
                pr.append((paese, score, a))
    pr.sort(key=lambda x: (-x[2]["n"], x[0], x[1]))
    headers = ["Paese", "Risultato_segnale", "N", "Altro_gol_n", "P_altro_gol", "IC95_low", "IC95_high", "E_gol_dopo"]
    data = [[p, s, a["n"], a["altro_gol"], a["p_altro_gol"], a["ic95_low"], a["ic95_high"], a["e_gol_dopo"]] for p, s, a in pr]
    ws = wb.create_sheet("Per_paese_risultato")
    write_sheet(ws, headers, data, pct_cols={5, 6, 7}, num_cols={8})

    headers = [
        "Paese",
        "N",
        "Lambda_min_tardi",
        *[item for m in P_FROM_MINUTES for item in (f"P_{m}_90", f"P_{m}_90plus3")],
    ]
    data = []
    for paese, pdata in sorted(lookup["paesi"].items(), key=lambda x: -x[1]["n"]):
        if pdata["n"] < 5:
            continue
        row = [paese, pdata["n"], pdata["lambda_min_tardi"]]
        for m in P_FROM_MINUTES:
            block = pdata["p_da_minuto"][str(m)]
            row.extend([block["p_90"], block["p_90_3"]])
        data.append(row)
    # global first
    g = lookup["globale"]
    grow = ["TUTTI", g["n"], g["lambda_min_tardi"]]
    for m in P_FROM_MINUTES:
        block = g["p_da_minuto"][str(m)]
        grow.extend([block["p_90"], block["p_90_3"]])
    data.insert(0, grow)
    pct_cols = set(range(4, 4 + 2 * len(P_FROM_MINUTES)))
    ws = wb.create_sheet("P_da_minuto")
    write_sheet(ws, headers, data, pct_cols=pct_cols, num_cols={3})

    camp_items = [(c, a) for c, a in lookup["campionati"].items()]
    camp_items.sort(key=lambda x: -x[1]["n"])
    headers = ["Campionato", "Paese", "N", "Altro_gol_n", "P_altro_gol", "IC95_low", "IC95_high", "E_gol_dopo", "Avg_minuto"]
    data = [
        [c, a["paese"], a["n"], a["altro_gol"], a["p_altro_gol"], a["ic95_low"], a["ic95_high"], a["e_gol_dopo"], a["avg_minuto"]]
        for c, a in camp_items
    ]
    ws = wb.create_sheet("Per_campionato")
    write_sheet(ws, headers, data, pct_cols={5, 6, 7}, num_cols={8, 9})

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)


def write_csv(rows: list[dict]) -> None:
    fields = [
        "id",
        "paese",
        "campionato",
        "casa",
        "ospite",
        "ht",
        "segnale_minuto",
        "segnale_risultato",
        "segnale_totale",
        "segnale_scarto",
        "segnale_stato",
        "ft",
        "gol_dopo",
        "altro_gol",
        "casa_dopo",
        "ospite_dopo",
        "minuti_residui_90",
    ]
    OUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUT_CSV.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader()
        for r in rows:
            row = {k: r[k] for k in fields}
            row["altro_gol"] = "SI" if r["altro_gol"] else "NO"
            w.writerow(row)


def write_html() -> None:
    html = r"""<!DOCTYPE html>
<html lang="it">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>LivePick — Segnali di ingresso</title>
<style>
:root { --bg:#0f172a; --card:#1e293b; --line:#334155; --txt:#e2e8f0; --mut:#94a3b8; --yes:#86efac; --no:#fca5a5; --acc:#38bdf8; }
* { box-sizing:border-box; margin:0; padding:0; }
body { font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif; background:var(--bg); color:var(--txt); padding:20px; }
h1 { font-size:22px; margin-bottom:6px; }
.sub { color:var(--mut); font-size:13px; margin-bottom:16px; }
.filters { display:grid; grid-template-columns:repeat(auto-fit,minmax(160px,1fr)); gap:10px; background:var(--card); padding:14px; border-radius:10px; margin-bottom:16px; }
.filters label { font-size:11px; color:var(--mut); display:block; margin-bottom:4px; }
.filters input, .filters select { width:100%; padding:8px 10px; border-radius:6px; border:1px solid var(--line); background:#0f172a; color:var(--txt); }
.stats { display:grid; grid-template-columns:repeat(auto-fit,minmax(130px,1fr)); gap:10px; margin-bottom:16px; }
.stat { background:var(--card); border-radius:10px; padding:12px; text-align:center; }
.stat .val { font-size:20px; font-weight:700; }
.stat .lbl { font-size:11px; color:var(--mut); margin-top:2px; }
.green { color:#22c55e; } .red { color:#ef4444; } .blue { color:#38bdf8; } .yellow { color:#f59e0b; }
.table-wrap { overflow:auto; background:var(--card); border-radius:10px; max-height:70vh; }
table { width:100%; border-collapse:collapse; font-size:13px; }
th { position:sticky; top:0; background:#334155; padding:8px; text-align:left; }
td { padding:7px 8px; border-bottom:1px solid var(--line); white-space:nowrap; }
tr:hover { background:#1a2744; }
.tag-yes { background:#166534; color:var(--yes); padding:1px 6px; border-radius:4px; font-size:11px; font-weight:700; }
.tag-no { background:#7f1d1d; color:var(--no); padding:1px 6px; border-radius:4px; font-size:11px; font-weight:700; }
a { color:var(--acc); }
.links { margin:10px 0 16px; font-size:13px; }
</style>
</head>
<body>
<h1>LivePick — segnali di ingresso</h1>
<div class="sub" id="sub">Caricamento...</div>
<div class="links">
  <a href="livepick_segnali.xlsx">Scarica Excel</a>
  · <a href="livepick_lookup.json">Lookup JSON</a>
</div>
<div class="filters">
  <div><label>Paese</label><select id="paese"><option value="">Tutti</option></select></div>
  <div><label>Campionato</label><select id="campionato"><option value="">Tutti</option></select></div>
  <div><label>Minuto segnale</label><input id="minuto" type="number" min="1" max="90" placeholder="es. 64"></div>
  <div><label>Risultato segnale</label><input id="risultato" placeholder="es. 2-1"></div>
  <div><label>Altro gol</label><select id="altro"><option value="">Tutti</option><option value="1">SI</option><option value="0">NO</option></select></div>
  <div><label>Cerca squadra</label><input id="squadra" placeholder="nome squadra"></div>
</div>
<div class="stats" id="stats"></div>
<div class="table-wrap">
<table>
<thead>
<tr>
<th>Paese</th><th>Campionato</th><th>Casa</th><th>Ospite</th><th>Min</th>
<th>Segnale</th><th>HT</th><th>FT</th><th>Gol dopo</th><th>Altro gol</th>
</tr>
</thead>
<tbody id="body"></tbody>
</table>
</div>
<script>
let rows = [];
const $ = id => document.getElementById(id);
function normScore(s){ return (s||'').replace(/\s+/g,'').replace('–','-'); }
function render(){
  const paese = $('paese').value;
  const camp = $('campionato').value;
  const minuto = $('minuto').value;
  const ris = normScore($('risultato').value);
  const altro = $('altro').value;
  const sq = $('squadra').value.trim().toLowerCase();
  const xs = rows.filter(r => {
    if (paese && r.paese !== paese) return false;
    if (camp && r.campionato !== camp) return false;
    if (minuto !== '' && String(r.segnale_minuto) !== String(parseInt(minuto,10))) return false;
    if (ris && r.segnale_risultato !== ris) return false;
    if (altro === '1' && !r.altro_gol) return false;
    if (altro === '0' && r.altro_gol) return false;
    if (sq && !(`${r.casa} ${r.ospite}`.toLowerCase().includes(sq))) return false;
    return true;
  });
  const n = xs.length;
  const k = xs.filter(r => r.altro_gol).length;
  const p = n ? k/n : 0;
  const e = n ? xs.reduce((s,r)=>s+r.gol_dopo,0)/n : 0;
  $('stats').innerHTML = `
    <div class="stat"><div class="val blue">${n}</div><div class="lbl">Partite filtrate</div></div>
    <div class="stat"><div class="val green">${k}</div><div class="lbl">Altro gol SI</div></div>
    <div class="stat"><div class="val red">${n-k}</div><div class="lbl">Altro gol NO</div></div>
    <div class="stat"><div class="val ${p>=0.7?'green':p>=0.5?'yellow':'red'}">${(p*100).toFixed(1)}%</div><div class="lbl">P altro gol dal segnale</div></div>
    <div class="stat"><div class="val yellow">${e.toFixed(2)}</div><div class="lbl">Gol medi dopo</div></div>`;
  const show = xs.slice(0, 500);
  $('body').innerHTML = show.map(r => `<tr>
    <td>${r.paese}</td><td>${r.campionato}</td><td>${r.casa}</td><td>${r.ospite}</td>
    <td>${r.segnale_minuto}</td><td>${r.segnale_risultato}</td><td>${r.ht}</td><td>${r.ft}</td>
    <td>${r.gol_dopo}</td>
    <td><span class="${r.altro_gol?'tag-yes':'tag-no'}">${r.altro_gol?'SI':'NO'}</span></td>
  </tr>`).join('') + (xs.length>500?`<tr><td colspan="10">Mostrate 500 di ${xs.length} righe. Restringi i filtri.</td></tr>`:'');
}
fetch('livepick_partite.json?'+Date.now()).then(r=>r.json()).then(data=>{
  rows = data;
  const paesi = [...new Set(rows.map(r=>r.paese))].sort();
  const camp = [...new Set(rows.map(r=>r.campionato))].sort();
  $('paese').innerHTML += paesi.map(p=>`<option>${p}</option>`).join('');
  $('campionato').innerHTML += camp.map(p=>`<option>${p}</option>`).join('');
  const q = new URLSearchParams(location.search);
  ['paese','campionato','minuto','risultato','altro','squadra'].forEach(id => {
    if (q.get(id)) $(id).value = q.get(id);
    $(id).addEventListener('input', render);
  });
  $('sub').textContent = `${rows.length} segnali storici · filtra paese / minuto / risultato per la probabilità di un altro gol dal segnale`;
  render();
}).catch(e => { $('sub').textContent = 'Errore: '+e.message; });
</script>
</body>
</html>
"""
    (ROOT / "docs" / "livepick.html").write_text(html, encoding="utf-8")


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Sorgente mancante: {SRC}")
    rows = load_rows()
    lookup = build_lookup(rows)
    write_csv(rows)
    OUT_PARTITE.write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
    OUT_LOOKUP.write_text(json.dumps(lookup, ensure_ascii=False), encoding="utf-8")
    build_xlsx(rows, lookup)
    print("sito=docs/livepick.html (non sovrascritto)")
    print(f"partite={len(rows)}")
    print(f"xlsx={OUT_XLSX}")
    print(f"lookup paesi={len(lookup['paesi'])}")
    india = lookup["paesi"].get("India", {})
    print("India n", india.get("n"), "p", india.get("p_altro_gol"))
    print("India 64", india.get("per_minuto", {}).get("64"))
    print("India 64|2-1", india.get("per_minuto_risultato", {}).get("64|2-1"))


if __name__ == "__main__":
    main()
